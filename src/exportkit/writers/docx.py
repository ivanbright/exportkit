"""Word (.docx) writer built on python-docx."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from ..exceptions import MissingDependencyError

_ALIGN = {"left": "LEFT", "center": "CENTER", "right": "RIGHT"}


def write_docx(report, raws: list[list[Any]], displays: list[list[str]], options: dict) -> bytes:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor
    except ImportError as exc:  # pragma: no cover - depends on install extras
        raise MissingDependencyError(
            "Word export needs python-docx. Install it with: pip install exportkit[docx]"
        ) from exc

    accent = str(options.get("accent_color", "242426")).lstrip("#")
    document = Document()
    document.add_heading(report.title, level=0)

    if report.subtitle:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(report.subtitle)
        run.italic = True

    for key, value in report.meta.items():
        paragraph = document.add_paragraph()
        label = paragraph.add_run(f"{key}: ")
        label.bold = True
        paragraph.add_run(str(value))

    table = document.add_table(rows=1, cols=len(report.columns))
    table.style = "Table Grid"

    def shade(cell, fill: str) -> None:
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), fill)
        properties.append(shading)

    def set_align(cell, value: str) -> None:
        cell.paragraphs[0].alignment = getattr(WD_ALIGN_PARAGRAPH, _ALIGN.get(value, "LEFT"))

    white = RGBColor(0xFF, 0xFF, 0xFF)
    header_cells = table.rows[0].cells
    for index, column in enumerate(report.columns):
        cell = header_cells[index]
        run = cell.paragraphs[0].add_run(column.header)
        run.bold = True
        run.font.color.rgb = white
        shade(cell, accent)
        set_align(cell, column.align)

    for display in displays:
        cells = table.add_row().cells
        for index, column in enumerate(report.columns):
            cells[index].text = display[index]
            set_align(cells[index], column.align)

    if report.summary:
        document.add_paragraph()
        for key, value in report.summary.items():
            paragraph = document.add_paragraph()
            label = paragraph.add_run(f"{key}: ")
            label.bold = True
            paragraph.add_run(str(value))

    note = report.footer or ""
    generated = report.created_at.strftime("%Y-%m-%d %H:%M")
    if note or generated:
        document.add_paragraph()
        paragraph = document.add_paragraph()
        run = paragraph.add_run(" · ".join(part for part in [note, f"Generated {generated}"] if part))
        run.italic = True
        run.font.size = Pt(9)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
