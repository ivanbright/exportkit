"""Excel (.xlsx) writer built on openpyxl."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from ..exceptions import MissingDependencyError

_INVALID_SHEET_CHARS = set('[]:*?/\\')


def _safe_sheet_name(name: str) -> str:
    cleaned = "".join(ch for ch in (name or "Report") if ch not in _INVALID_SHEET_CHARS)
    return (cleaned or "Report")[:31]


def write_xlsx(report, raws: list[list[Any]], displays: list[list[str]], options: dict) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - depends on install extras
        raise MissingDependencyError(
            "Excel export needs openpyxl. Install it with: pip install exportkit[xlsx]"
        ) from exc

    accent = str(options.get("accent_color", "242426")).lstrip("#")
    header_fill = PatternFill("solid", fgColor=accent)
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14)
    subtitle_font = Font(italic=True, size=11, color="4D4F57")
    label_font = Font(bold=True)
    bold = Font(bold=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _safe_sheet_name(report.sheet_name)

    columns = report.columns
    ncols = len(columns)
    cursor = 1

    def merged_row(text: str, font, height: int = 20) -> None:
        nonlocal cursor
        sheet.cell(row=cursor, column=1, value=text).font = font
        if ncols > 1:
            sheet.merge_cells(start_row=cursor, start_column=1, end_row=cursor, end_column=ncols)
        sheet.row_dimensions[cursor].height = height
        cursor += 1

    merged_row(report.title, title_font, height=24)
    if report.subtitle:
        merged_row(report.subtitle, subtitle_font)
    for key, value in report.meta.items():
        sheet.cell(row=cursor, column=1, value=f"{key}:").font = label_font
        sheet.cell(row=cursor, column=2, value=str(value))
        cursor += 1

    cursor += 1  # spacer

    header_row = cursor
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=header_row, column=index, value=column.header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal=column.align, vertical="center")
    cursor += 1

    for raw in raws:
        for index, (column, value) in enumerate(zip(columns, raw), start=1):
            cell = sheet.cell(row=cursor, column=index, value=value)
            cell.alignment = Alignment(horizontal=column.align, vertical="top")
        cursor += 1

    if raws:
        last_data_row = cursor - 1
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{last_data_row}"

    if report.summary:
        cursor += 1
        for key, value in report.summary.items():
            sheet.cell(row=cursor, column=1, value=f"{key}:").font = bold
            sheet.cell(row=cursor, column=2, value=str(value))
            cursor += 1

    # column widths from the widest display value
    for index, column in enumerate(columns, start=1):
        if column.width:
            sheet.column_dimensions[get_column_letter(index)].width = column.width
            continue
        longest = len(column.header)
        for display in displays:
            longest = max(longest, len(display[index - 1]))
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 60)

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
