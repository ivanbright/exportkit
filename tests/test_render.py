import io
import sys

import pytest

from exportkit import Column, Format, Report, render, render_to
from exportkit.exceptions import MissingDependencyError


@pytest.fixture
def report():
    return Report(
        title="Waiting list",
        subtitle="I&M Bank \u00b7 Nairobi CBD",
        meta={"Service": "Teller"},
        columns=[
            Column("Ticket", align="right"),
            Column("Customer"),
            Column("Wait", align="right"),
        ],
        rows=[[1201, "Amina Otieno", "~6 min"], [1202, "Brian Kamau", "now"]],
        summary={"Total": 2},
        footer="Don't Queue",
    )


def test_xlsx_round_trip(report):
    from openpyxl import load_workbook

    data = render(report, "xlsx")
    assert data[:2] == b"PK"

    sheet = load_workbook(io.BytesIO(data)).active
    assert sheet["A1"].value == "Waiting list"
    values = {
        str(cell.value)
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    }
    assert "Ticket" in values
    assert "Amina Otieno" in values
    # raw numeric values survive as numbers, not text
    assert any(
        cell.value == 1201 for row in sheet.iter_rows() for cell in row
    )


def test_xlsx_accepts_aware_datetimes():
    from datetime import datetime, timezone

    from openpyxl import load_workbook

    aware = datetime(2026, 10, 14, 9, 30, tzinfo=timezone.utc)
    report = Report(title="Times", columns=[Column("When")], rows=[[aware]])

    data = render(report, "xlsx")
    sheet = load_workbook(io.BytesIO(data)).active
    stored = [
        cell.value
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, datetime)
    ]
    assert stored, "datetime cell was not written"
    assert stored[0].tzinfo is None


def test_docx_round_trip(report):
    from docx import Document

    data = render(report, "docx")
    assert data[:2] == b"PK"

    document = Document(io.BytesIO(data))
    assert document.paragraphs[0].text == "Waiting list"
    table = document.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Ticket", "Customer", "Wait"]
    assert table.rows[1].cells[1].text == "Amina Otieno"
    assert table.rows[2].cells[0].text == "1202"


def test_pdf_round_trip(report):
    data = render(report, Format.PDF)
    assert data.startswith(b"%PDF")
    assert data.rstrip().endswith(b"%%EOF")
    assert len(data) > 500


def test_render_to_stream(report):
    buffer = io.BytesIO()
    render_to(report, "xlsx", buffer)
    assert buffer.getvalue()[:2] == b"PK"


def test_render_to_path(report, tmp_path):
    target = tmp_path / "out.pdf"
    result = render_to(report, "pdf", target)
    assert result == str(target)
    assert target.read_bytes().startswith(b"%PDF")


def test_missing_dependency_message(report, monkeypatch):
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    with pytest.raises(MissingDependencyError) as excinfo:
        render(report, "xlsx")
    assert "exportkit[xlsx]" in str(excinfo.value)
